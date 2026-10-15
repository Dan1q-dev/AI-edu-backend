import ctypes
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from django.test import SimpleTestCase
from .importing.document import to_blocks
from .importing.parser import Extractor


class ImportFormattingTests(SimpleTestCase):
    def extract_nodes(self, document, extension, folder):
        source = Path(folder) / ('source.' + extension)
        document.save(source)
        result = Extractor(folder, {'max_pages': 100}).extract(source)
        return json.loads(to_blocks(result['elements'], {})[0]['content'])['content']

    def test_word_partial_run_colors_highlight_theme_and_table_formatting(self):
        from docx import Document
        from docx.shared import RGBColor
        from docx.enum.text import WD_COLOR_INDEX
        from docx.enum.dml import MSO_THEME_COLOR_INDEX
        document = Document()
        paragraph = document.add_paragraph('Plain ')
        colored = paragraph.add_run('red and yellow')
        colored.bold = True
        colored.font.color.rgb = RGBColor.from_string('DC2626')
        colored.font.highlight_color = WD_COLOR_INDEX.YELLOW
        paragraph.add_run(' plain again')
        theme = document.add_paragraph().add_run('Theme accent')
        theme.font.color.theme_color = MSO_THEME_COLOR_INDEX.ACCENT_1
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].add_run('Table color').font.color.rgb = RGBColor.from_string('2563EB')
        with TemporaryDirectory() as folder:
            nodes = self.extract_nodes(document, 'docx', folder)
        self.assertNotIn('marks', nodes[0]['content'][0])
        self.assertEqual(nodes[0]['content'][1], {'type': 'text', 'text': 'red and yellow', 'marks': [
            {'type': 'bold'}, {'type': 'textStyle', 'attrs': {'color': '#dc2626'}}, {'type': 'highlight', 'attrs': {'color': '#ffff00'}}]})
        self.assertNotIn('marks', nodes[0]['content'][2])
        self.assertEqual(nodes[1]['content'][0]['marks'], [{'type': 'textStyle', 'attrs': {'color': '#4f81bd'}}])
        self.assertEqual(nodes[2]['content'][0]['content'][0]['content'][0]['content'][0]['marks'], [{'type': 'textStyle', 'attrs': {'color': '#2563eb'}}])

    def test_word_direct_numbering_preserves_nesting_restart_and_continuation(self):
        from docx import Document
        from docx.oxml import parse_xml, OxmlElement
        from docx.oxml.ns import qn, nsdecls
        document = Document()
        root = document.part.numbering_part.element
        root.append(parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="42"><w:lvl w:ilvl="0"><w:start w:val="4"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl><w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/></w:lvl></w:abstractNum>'))
        root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="77"><w:abstractNumId w:val="42"/></w:num>'))
        root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="78"><w:abstractNumId w:val="42"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
        def item(text, num_id='77', level=0):
            paragraph = document.add_paragraph(text)
            props = paragraph._p.get_or_add_pPr()
            numbering = OxmlElement('w:numPr')
            for tag, value in [('ilvl', str(level)), ('numId', num_id)]:
                child = OxmlElement('w:' + tag); child.set(qn('w:val'), value); numbering.append(child)
            props.append(numbering)
        item('Fourth')
        item('Nested bullet', level=1)
        item('Fifth')
        item('New sequence', '78')
        document.add_paragraph('Body paragraph')
        item('Sixth')
        with TemporaryDirectory() as folder:
            nodes = self.extract_nodes(document, 'docx', folder)
        self.assertEqual([node['type'] for node in nodes], ['orderedList', 'orderedList', 'paragraph', 'orderedList'])
        self.assertEqual([nodes[index]['attrs']['start'] for index in (0, 1, 3)], [4, 1, 6])
        self.assertEqual(len(nodes[0]['content']), 2)
        self.assertEqual(nodes[0]['content'][0]['content'][1]['type'], 'bulletList')
        self.assertEqual(nodes[0]['content'][0]['content'][1]['content'][0]['content'][0]['content'][0]['text'], 'Nested bullet')

    def test_presentation_colors_highlight_and_numbered_lists(self):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches
        from pptx.oxml.xmlchemy import OxmlElement
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3)).text_frame
        first = frame.paragraphs[0]
        first.text = 'Fourth colored item'
        first.runs[0].font.color.rgb = RGBColor.from_string('DC2626')
        highlight = OxmlElement('a:highlight'); rgb = OxmlElement('a:srgbClr'); rgb.set('val', 'FEF08A'); highlight.append(rgb)
        first.runs[0]._r.get_or_add_rPr().append(highlight)
        numbering = OxmlElement('a:buAutoNum'); numbering.set('type', 'arabicPeriod'); numbering.set('startAt', '4')
        first._p.get_or_add_pPr().append(numbering)
        nested = frame.add_paragraph(); nested.text = 'Nested bullet'; nested.level = 1
        bullet = OxmlElement('a:buChar'); bullet.set('char', '•'); nested._p.get_or_add_pPr().append(bullet)
        second = frame.add_paragraph(); second.text = 'Fifth item'
        numbering = OxmlElement('a:buAutoNum'); numbering.set('type', 'arabicPeriod'); second._p.get_or_add_pPr().append(numbering)
        with TemporaryDirectory() as folder:
            nodes = self.extract_nodes(document, 'pptx', folder)
        self.assertEqual(len(nodes), 1)
        self.assertEqual(nodes[0]['attrs']['start'], 4)
        self.assertEqual(len(nodes[0]['content']), 2)
        self.assertEqual(nodes[0]['content'][0]['content'][0]['content'][0]['marks'], [
            {'type': 'textStyle', 'attrs': {'color': '#dc2626'}}, {'type': 'highlight', 'attrs': {'color': '#fef08a'}}])
        self.assertEqual(nodes[0]['content'][0]['content'][1]['type'], 'bulletList')

    def test_presentation_inherited_placeholder_bullets(self):
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = 'Slide title'
        frame = slide.placeholders[1].text_frame
        frame.text = 'First bullet'
        nested = frame.add_paragraph(); nested.text = 'Nested bullet'; nested.level = 1
        with TemporaryDirectory() as folder:
            nodes = self.extract_nodes(document, 'pptx', folder)
        self.assertEqual([node['type'] for node in nodes], ['heading', 'bulletList'])
        self.assertEqual(nodes[1]['content'][0]['content'][1]['type'], 'bulletList')

    def test_native_pdf_colors_and_highlight_annotations_match_source_glyphs(self):
        import pypdfium2 as pdfium
        from .test_import_ocr import make_pdf
        from .importing.pdf_style import enrich_pdf_colors
        with TemporaryDirectory() as folder:
            source, annotated = Path(folder) / 'source.pdf', Path(folder) / 'annotated.pdf'
            make_pdf(source, stream=b'BT /F1 24 Tf 50 700 Td 1 0 0 rg (Red ) Tj 0 0 1 rg (blue) Tj ET')
            with pdfium.PdfDocument(source) as document:
                page = document[0]
                annot = pdfium.raw.FPDFPage_CreateAnnot(page, pdfium.raw.FPDF_ANNOT_HIGHLIGHT)
                rect = pdfium.raw.FS_RECTF(45, 730, 200, 690)
                pdfium.raw.FPDFAnnot_SetRect(annot, ctypes.byref(rect))
                pdfium.raw.FPDFAnnot_SetColor(annot, pdfium.raw.FPDFANNOT_COLORTYPE_Color, 255, 255, 0, 255)
                pdfium.raw.FPDFPage_CloseAnnot(annot)
                document.save(annotated)
                page.close()
            bbox = {'l': 45, 'b': 690, 'r': 200, 't': 730, 'coord_origin': 'BOTTOMLEFT'}
            elements = [{'kind': 'paragraph', 'text': 'Red blue', 'page': 1, 'bbox': bbox},
                        {'kind': 'paragraph', 'text': 'Red blue', 'page': 1, 'bbox': bbox, 'ocr': True},
                        {'kind': 'paragraph', 'text': 'Different text', 'page': 1, 'bbox': bbox}]
            enrich_pdf_colors(annotated, elements)
        self.assertEqual(''.join(node['text'] for node in elements[0]['inline']), 'Red blue')
        self.assertEqual(elements[0]['inline'][0]['marks'], [{'type': 'textStyle', 'attrs': {'color': '#ff0000'}}, {'type': 'highlight', 'attrs': {'color': '#ffff00'}}])
        self.assertEqual(elements[0]['inline'][1]['marks'][0]['attrs']['color'], '#0000ff')
        self.assertNotIn('inline', elements[1])
        self.assertNotIn('inline', elements[2])
