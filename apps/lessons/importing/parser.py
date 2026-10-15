"""Native Office structure + layout-aware Docling PDF/OCR, executed only in worker child."""
import io
import json
import subprocess
from pathlib import Path
from PIL import Image
from .document import ImportFailure


class Extractor:
    def __init__(self, output, options):
        self.output, self.options = Path(output), options
        self.elements, self.assets, self.warnings = [], [], []
        self.characters = 0

    def add(self, kind, text='', **extra):
        if not text and kind not in ('image', 'table'):
            return
        size = len(text) + sum(len(cell.get('text', '')) for row in extra.get('rows', []) for cell in row)
        if len(self.elements) >= 5000 or self.characters + size > 500000:
            raise ImportFailure('document_limit')
        self.characters += size
        # Hyperlinks, breaks and fields may not be represented by paragraph.runs.
        # Fall back to complete plain text rather than losing those fragments.
        if 'inline' in extra and ''.join(node.get('text', '') for node in extra['inline']) != text:
            extra.pop('inline')
        self.elements.append({'source_id': f'e{len(self.elements) + 1}', 'kind': kind, 'text': text, **extra})

    def image(self, data, page=None, bbox=None, ocr=True):
        try:
            with Image.open(io.BytesIO(data)) as original:
                image = original.convert('RGB')
                image.thumbnail((2400, 2400))
                key = f'image-{len(self.assets)}'
                path = self.output / f'{key}.png'
                image.save(path, 'PNG')
        except (OSError, ValueError):
            self.warnings.append('unsupported_image')
            return
        if len(self.assets) >= 150:
            raise ImportFailure('document_limit')
        self.assets.append({'key': key, 'file': path.name, 'page': page, 'is_page': False})
        self.add('image', image=key, page=page, bbox=bbox)
        if ocr:
            try:
                result = subprocess.run(['tesseract', str(path), 'stdout', '-l', self.options['languages'], '--psm', '3'],
                                        capture_output=True, timeout=60, encoding='utf-8')
            except FileNotFoundError:
                raise ImportFailure('ocr_unavailable') from None
            except subprocess.TimeoutExpired:
                raise ImportFailure('timeout') from None
            if result.returncode:
                raise ImportFailure('ocr_unavailable')
            if result.stdout.strip():
                self.add('paragraph', result.stdout.strip(), page=page, bbox=bbox, ocr=True)
                self.warnings.append('image_ocr')

    def office_runs(self, runs):
        inline = []
        for run in runs:
            text = run.text
            if not text:
                continue
            marks = [{'type': name} for name in ('bold', 'italic', 'underline') if getattr(run, name, None) or (hasattr(run, 'font') and getattr(run.font, name, None))]
            node = {'type': 'text', 'text': text}
            if marks:
                node['marks'] = marks
            inline.append(node)
        return inline

    def docx(self, path):
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn
        document = Document(path)

        def images(element):
            for blip in element.xpath('.//a:blip'):
                relation = document.part.related_parts.get(blip.get(qn('r:embed')))
                if relation:
                    self.image(relation.blob)

        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text
                style = (block.style.name if block.style else '').lower()
                kind = 'paragraph'
                level = 2
                if style.startswith(('heading', 'title', 'заголовок')):
                    kind = 'heading'; level = 3 if any(x in style for x in ('2', '3', '4', '5', '6')) else 2
                elif 'quote' in style or 'цитат' in style:
                    kind = 'quote'
                elif 'code' in style or 'код' in style:
                    kind = 'code'
                elif 'list bullet' in style:
                    kind = 'bullet'
                elif 'list number' in style:
                    kind = 'ordered'
                self.add(kind, text, inline=self.office_runs(block.runs), level=level, page=None)
                math = block._p.xpath('.//m:t')
                if math:
                    self.add('paragraph', ''.join(node.text or '' for node in math), page=None)
                    self.warnings.append('native_formula')
                images(block._p)
            elif isinstance(block, Table):
                rows = []
                # Expand Word merges into a rectangular table to retain all cell text.
                seen = set()
                for row in block.rows:
                    cells = []
                    for cell in row.cells:
                        text = cell.text if cell._tc not in seen else ''
                        seen.add(cell._tc)
                        cells.append({'text': text})
                    rows.append(cells)
                self.add('table', rows=rows, page=None)
                images(block._tbl)
        self.warnings.append('office_layout')
        # Header/footer material is retained explicitly, without dropping it silently.
        seen_parts = set()
        for section in document.sections:
            for part in (section.header, section.footer):
                if part.part.partname in seen_parts:
                    continue
                seen_parts.add(part.part.partname)
                for paragraph in part.paragraphs:
                    if paragraph.text.strip():
                        self.add('paragraph', paragraph.text, page=None)
                        self.warnings.append('header_footer')

    def pptx(self, path):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        presentation = Presentation(path)
        if len(presentation.slides) > self.options['max_pages']:
            raise ImportFailure('document_limit')

        def shapes(items, page, title):
            for shape in sorted(items, key=lambda x: (x.top, x.left)):
                bbox = {'left': shape.left, 'top': shape.top, 'width': shape.width, 'height': shape.height, 'unit': 'emu'}
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes(shape.shapes, page, title)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    self.image(shape.image.blob, page, bbox)
                elif shape.has_table:
                    rows = []
                    for row in shape.table.rows:
                        cells = []
                        for cell in row.cells:
                            if not cell.is_spanned:
                                cells.append({'text': cell.text, 'colspan': cell.span_width, 'rowspan': cell.span_height})
                        rows.append(cells)
                    self.add('table', rows=rows, page=page, bbox=bbox)
                elif shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        kind = 'heading' if shape == title else 'paragraph'
                        if kind == 'paragraph' and paragraph._p.xpath('./a:pPr/a:buChar'):
                            kind = 'bullet'
                        self.add(kind, paragraph.text, inline=self.office_runs(paragraph.runs), page=page, bbox=bbox)
                elif shape.has_chart:
                    self.warnings.append('chart_not_imported')
        for page, slide in enumerate(presentation.slides, 1):
            shapes(slide.shapes, page, slide.shapes.title)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes and notes.text.strip():
                    self.add('paragraph', notes.text, page=page)
        self.warnings.append('office_layout')

    def pdf(self, path):
        import pypdfium2 as pdfium
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions, TesseractCliOcrOptions, OcrMode
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling_core.types.doc import PictureItem, TableItem
        with pdfium.PdfDocument(path) as pdf:
            if len(pdf) > self.options['max_pages']:
                raise ImportFailure('document_limit')
        options = PdfPipelineOptions()
        options.do_ocr = True
        options.do_table_structure = True
        options.generate_picture_images = True
        options.generate_page_images = True
        options.images_scale = 1.5
        try:
            languages = subprocess.run(['tesseract', '--list-langs'], capture_output=True, encoding='utf-8', timeout=10)
        except (FileNotFoundError, subprocess.TimeoutExpired):
            raise ImportFailure('ocr_unavailable') from None
        if languages.returncode or not set(self.options['languages'].split('+')).issubset(set(languages.stdout.split())):
            raise ImportFailure('ocr_unavailable')
        options.ocr_options = TesseractCliOcrOptions(lang=self.options['languages'].split('+'), mode=OcrMode.FULL_PAGE if self.options['force_ocr'] else OcrMode.DEFAULT)
        converted = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)}).convert(path)
        if converted.status.value != 'success':
            raise ImportFailure('invalid_document')
        document = converted.document
        for page_no, page in document.pages.items():
            if page.image:
                key = f'page-{page_no}'
                page.image.pil_image.save(self.output / f'{key}.png', 'PNG')
                self.assets.append({'key': key, 'file': f'{key}.png', 'page': page_no, 'is_page': True})
        labels = {'title': 'heading', 'section_header': 'heading', 'code': 'code', 'formula': 'formula', 'list_item': 'bullet'}
        for item, _ in document.iterate_items():
            provenance = item.prov[0] if getattr(item, 'prov', []) else None
            extra = {'page': provenance.page_no if provenance else None,
                     'bbox': provenance.bbox.model_dump(mode='json') if provenance else None}
            if isinstance(item, PictureItem):
                image = item.get_image(document)
                if image:
                    data = io.BytesIO(); image.save(data, 'PNG')
                    self.image(data.getvalue(), **extra, ocr=False)
            elif isinstance(item, TableItem):
                data = item.data
                if data.num_rows > 200 or data.num_cols > 50:
                    raise ImportFailure('document_limit')
                grid = data.grid
                rows = [[] for _ in range(data.num_rows)]
                occupied = set()
                for row in range(data.num_rows):
                    for col in range(data.num_cols):
                        if (row, col) in occupied:
                            continue
                        cell = grid[row][col]
                        for r in range(cell.start_row_offset_idx, cell.end_row_offset_idx):
                            for c in range(cell.start_col_offset_idx, cell.end_col_offset_idx):
                                occupied.add((r, c))
                        rows[row].append({'text': cell.text, 'header': cell.column_header,
                                          'colspan': cell.col_span, 'rowspan': cell.row_span})
                self.add('table', rows=rows, **extra)
            elif hasattr(item, 'text'):
                label = item.label.value
                self.add(labels.get(label, 'paragraph'), item.text, **extra)
        self.warnings.append('ocr_review')

    def extract(self, path):
        try:
            {'.docx': self.docx, '.pptx': self.pptx, '.pdf': self.pdf}[Path(path).suffix.lower()](path)
        except ImportError:
            raise ImportFailure('dependencies_missing') from None
        except ImportFailure:
            raise
        except Exception:
            raise ImportFailure('invalid_document') from None
        if not self.elements:
            raise ImportFailure('empty_document')
        return {'elements': self.elements, 'assets': self.assets, 'warnings': list(dict.fromkeys(self.warnings))}


def process(config):
    output = Path(config['output'])
    def stage(value):
        (output / 'stage').write_text(value, encoding='utf-8')
    stage('parsing')
    result = Extractor(output, config).extract(config['source'])
    # Separate extraction and classification stages without returning raw HTML.
    (output / 'extracted.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    if config['use_ai']:
        stage('structuring')
        from .structure import classify
        result['warnings'].extend(classify(result['elements'], config['ai']))
    stage('validating')
    from .document import to_blocks
    to_blocks(result['elements'], {asset['key']: 1 for asset in result['assets']})
    (output / 'result.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')


if __name__ == '__main__':
    import sys
    config = json.load(sys.stdin)
    try:
        process(config)
    except ImportFailure as error:
        Path(config['output'], 'error').write_text(error.code, encoding='utf-8')
    except Exception:
        Path(config['output'], 'error').write_text('processing_failed', encoding='utf-8')
