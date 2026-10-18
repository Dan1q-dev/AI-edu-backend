"""Package already executed follow-up evidence; retain all 17 October experiments."""
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
from pathlib import Path
import platform
import re
import shutil
import subprocess
import sys
import zipfile

ROOT=Path(__file__).resolve().parents[2]
MATERIALS=ROOT/'AIEdu Supplementary Materials'
WORK=ROOT/'.validation-work'
STAMP='2026-10-18'
S1=MATERIALS/'supplementary/S1 configuration and source'/('Portal integration revision '+STAMP)
S3=MATERIALS/'supplementary/S3 browser CNN'/('Model verification '+STAMP)
S4=MATERIALS/'supplementary/S4 API and validation'
PORTAL=S4/('Portal integration validation '+STAMP)
RETENTION=S4/('Retention validation '+STAMP)


def read_json(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def write_json(p,value):p.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read_log(p):
    raw=p.read_bytes()
    return raw.decode('utf-16' if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig',errors='replace')


def main():
    portal=read_json(WORK/'portal-final/results.json');retention=read_json(WORK/'retention-final/results.json')
    frontend=read_json(WORK/'frontend-tests.json');cnn=read_json(WORK/'cnn-final/replay results.json')
    backend_log=read_log(WORK/'backend-tests.log')
    assert portal['failed']==0 and portal['passed']==17
    assert retention['failed']==0 and retention['passed']==9 and not retention['skipped']
    assert frontend['numFailedTests']==0 and frontend['numPassedTests']==142
    assert re.search(r'Ran 69 tests',backend_log) and re.search(r'\nOK\s*\n',backend_log)
    assert cnn['passed']
    for directory in (S1,S3,PORTAL,RETENTION):directory.mkdir(parents=True,exist_ok=False)
    # Preserve the original reader-facing description before updating it.
    shutil.copy2(MATERIALS/'AIEdu Supplementary Materials.docx',S1/'Original supplement description 17 October.docx')
    shutil.copy2(MATERIALS/'README.txt',S1/'Original README 17 October.txt')
    for source,target in [(WORK/'portal-final',PORTAL),(WORK/'retention-final',RETENTION),(WORK/'cnn-final',S3)]:
        for p in source.rglob('*'):
            if not p.is_file():continue
            dest=target/p.relative_to(source);dest.parent.mkdir(parents=True,exist_ok=True)
            if p.suffix=='.log':
                text=read_log(p).replace(str(ROOT),'<WORKSPACE>')
                text=re.sub(r'C:\\Users\\[^\\\s]+\\AppData\\Local\\Temp\\aiedu-[^\\\s"\']+', '<TEMPORARY_TEST_DIRECTORY>',text)
                dest.write_text(text,encoding='utf-8')
            else:shutil.copy2(p,dest)
    for name in ['backend-tests.log','frontend-tests.log','frontend-build.log','ai-typecheck.log','ai-build.log','domain-schema.log']:
        (PORTAL/name).write_text(read_log(WORK/name).replace(str(ROOT),'<WORKSPACE>'),encoding='utf-8')
    for name in ['frontend-tests.json','domain-current-openapi.yaml']:shutil.copy2(WORK/name,PORTAL/name)
    revision_sources=[]
    for kind,folder in [('backend',ROOT/'AI-edu-backend'),('frontend',ROOT/'AI-edu-frontend')]:
        commit=subprocess.check_output(['git','-C',str(folder),'rev-parse','HEAD']).decode().strip()
        names=subprocess.check_output(['git','-C',str(folder),'ls-files','--cached','--others','--exclude-standard','-z']).decode('utf-8').split('\0')
        files=[]
        for name in sorted(set(names)):
            if not name:continue
            p=folder/name
            if not p.is_file():continue
            if p.name.startswith('.env') and p.name!='.env.example':raise RuntimeError('Private environment file in source list')
            if p.suffix.lower() in ('.sqlite','.sqlite3','.db','.log','.pyc'):continue
            files.append(p)
        archive=S1/f'{kind} portal followup {commit[:12]}.zip'
        with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
            for p in files:z.write(p,p.relative_to(folder).as_posix())
        dirty=bool(subprocess.check_output(['git','-C',str(folder),'diff','--name-only']).strip())
        manifest={'artifact':kind+'_portal_followup','base_commit':commit,'working_tree_changes':dirty,
            'archive':str(archive.relative_to(MATERIALS/'supplementary/S1 configuration and source')).replace('\\','/'),
            'sha256':sha(archive),'purpose':'18 October integration/failure/retention follow-up; not the 17 October RAG configuration',
            'files':[{'file':p.relative_to(folder).as_posix(),'sha256':sha(p)} for p in files]}
        revision_sources.append(manifest)
    write_json(S1/'source provenance.json',revision_sources)
    patch=subprocess.check_output(['git','-C',str(ROOT/'AI-edu-backend'),'diff','--','TUTOR_INTEGRATION.md','services/ai-tutor/.env.example','services/ai-tutor/app/api/chat/route.ts'])
    (S1/'gateway followup.patch').write_bytes(patch)
    # Link new provenance without changing identifiers of the original archives.
    old_provenance=MATERIALS/'supplementary/S1 configuration and source/repository provenance.json'
    provenance=read_json(old_provenance)
    provenance.extend({k:v for k,v in item.items() if k!='files'} for item in revision_sources)
    write_json(old_provenance,provenance)
    provider_versions={name:read_json(ROOT/'AI-edu-backend/services/ai-tutor/node_modules'/name/'package.json')['version']
        for name in ['next','openai','ioredis','@prisma/client','typescript']}
    runtime={'date_utc':datetime.now(timezone.utc).isoformat(),'platform':platform.platform(),'python':platform.python_version(),
        'node':subprocess.check_output(['node','--version']).decode().strip(),'ai_packages':provider_versions,
        'python_packages':{n:importlib.metadata.version(n) for n in ['Django','djangorestframework','numpy','python-docx']},
        'redis_test_server':'Redis 5.0.14.1 Windows port, isolated process; not the production Redis 7 deployment',
        'redis_download':{'source':'https://github.com/tporadowski/redis/releases/download/v5.0.14.1/Redis-x64-5.0.14.1.zip',
            'archive_sha256':'018ea18a35876383cbb5f4cd0258adfc87747cf9d619bce1cf73a2e36f720ccf'},
        'postgres_scope':retention['postgres_schema_scope'],'provider':'local HTTP fixture, synthetic-validation-model; no live OpenAI request'}
    write_json(S1/'runtime versions.json',runtime)
    write_json(PORTAL/'source references.json',[{k:v for k,v in s.items() if k!='files'} for s in revision_sources])
    write_json(PORTAL/'regression summary.json',{'backend_tests':{'passed':69,'failed':0,'command':'python manage.py test apps.tutor apps.education apps.assessments --noinput -v 2'},
        'frontend_tests':{'passed':142,'failed':0,'command':'npm test -- --reporter=json'},
        'frontend_build':{'exit_code':0,'command':'npm run build'},'ai_typecheck':{'exit_code':0,'command':'tsc --noEmit --incremental false'},
        'ai_build':{'exit_code':0,'command':'next build','environment':'disposable source copy, synthetic provider key'},
        'domain_schema':{'exit_code':0,'command':'spectacular --validate','database':'isolated SQLite :memory:'}})
    write_json(RETENTION/'deployment confirmation status.json',{
        'provider_request_store_flag':{'status':'VERIFIED_IN_LOCAL_HTTP_TEST','value':False,'evidence':'../Portal integration validation 2026-10-18/provider requests.json'},
        'provider_account_retention':{'status':'NOT_VERIFIED','reason':'No provider account access; no real key used'},
        'production_environment':{'status':'NOT_VERIFIED','reason':'Only disposable local services were started'},
        'production_deletion':{'status':'NOT_VERIFIED','required_operator_evidence':['Deployed retention settings with secrets removed','Timestamped cleanup invocation and before/after record counts','Backup and replica retention policy if those copies are in scope']},
        'production_redis':{'status':'NOT_VERIFIED','required_operator_evidence':['Whether Redis is enabled in the deployed gateway','Redis server version and observed TTL/PTTL for each actual key category']},
        'verified_local_deletion':{'databases':['SQLite','PostgreSQL'],'retention_days':30,'maintenance_interval_requests':25,
            'trigger':'application activity, not scheduled wall-clock expiry','expired_and_fresh_fixture_records_checked':True},
        'verified_local_redis_ttl':{'quota_minute_seconds':120,'quota_day_seconds':259200,'quota_month_seconds':3888000,
            'quota_reservation_seconds':600,'job_status_seconds':1209600,'accelerated_expiry_test_seconds':1,
            'not_dialogue_history_ttl':True}})
    (RETENTION/'README.txt').write_text('''Retention evidence, 18 October 2026

results.json combines nine passing checks against new local test instances.
The SQLite and PostgreSQL checks seed expired and fresh records in each of four
tables, confirm no cleanup before request 25, then confirm only expired records
are removed by the actual application maintenance function. This is an activity
trigger, not a guarantee of deletion exactly 30 days after insertion.

PostgreSQL retention tables use the repository schema unchanged. The optional
vector column/index was omitted because pgvector is not installed on this host;
vector retrieval is not tested here. The 17 October pgvector experiment is separate.

Redis checks observe actual TTLs on five key categories using default settings,
then use explicitly overridden one-second TTLs to verify physical key expiration.
Redis is a fresh loopback-only 5.0.14.1 Windows test process. Production Compose
references Redis 7; no production Redis compatibility or TTL attestation is made.

deployment confirmation status.json intentionally leaves provider-account and
production-operation items NOT_VERIFIED. The observed store=false flag is request
evidence; it does not certify provider-side deletion or account-level retention.
No production database, Redis data, provider account or running deployment was changed.
Reproduce using validation/retention_runner.py in the new backend source archive.
''',encoding='utf-8')
    (PORTAL/'README.txt').write_text('''Portal integration evidence, 18 October 2026

results.json contains 17 passing checks. Chrome logs in two synthetic learners
using the real portal form. Their requests pass through Vite, Django session/CSRF
authentication and lesson permissions, the real Next.js route, OpenAI SDK, lesson
retrieval and SQLite persistence. Only the external provider is simulated over HTTP.
Embeddings and answers from this fixture are deterministic transport fixtures;
these results do not evaluate model quality or live provider connectivity.

Three multi-turn browser rounds execute concurrently. Provider timestamps prove
overlap of the first two response requests. Browser histories and persisted
gateway prompts are checked for cross-user/cross-lesson contamination. Source
text and identity are server-derived; client-supplied spoofed fields are ignored.
Anonymous and foreign-lecture requests are refused before reaching the gateway.

Six failure cases: embedding/retrieval outage, provider HTTP 503, provider header
timeout, clean EOF without response.completed, stalled response stream, malformed
provider SSE. Each produces a visible browser error and no completed assistant
record. A retrieval error maps to portal HTTP 503. Generation errors may arrive
as SSE event:error within HTTP 200. Chrome may discard the raw body after the real
client cancels its reader on error; these cases retain visible error and database
evidence instead of inventing a raw response body.

Fixes: explicit store=false, validated configurable provider timeout/retry count,
deadline for streamed generation, and rejection of streams lacking completion.
Provider timeout is 1500 ms and retries zero only for this controlled run.
Production defaults are 60000 ms and two retries. Error messages were repaired.

The new source references identify current frontend and backend snapshots. The
17 October S1 archives and S4 reports remain historical evidence for that date;
their 22/12 test counts do not describe the current source. The new regression
counts are 142 frontend and 69 backend tests. Build/typecheck/schema logs are included.

Reproduce: extract the new source archives beside one another, install their
existing lockfile dependencies, generate Prisma client, then run
python validation/portal_gateway.py --output <NEW_DIRECTORY> from the backend.
See validation/README.txt for prerequisites. Screenshots contain synthetic data.
Some log filesystem prefixes were replaced by <WORKSPACE>/<TEMPORARY_TEST_DIRECTORY>.
''',encoding='utf-8')
    (S3/'Model scope.txt').write_text('''Evaluated model selection and evidence scope

The evaluated object is the browser JavaScript ten-class (0-9) classifier with
294 fixed convolution coefficients and 6762 trainable dense parameters. It is
not represented as a PyTorch-trained model or as an export_viz.py export.

The 18 October replay loads the original archived cnn.js/mlp.js modules, frozen
checkpoint and all 200 retained pixel arrays. Every predicted class matches;
the maximum probability difference is recorded in replay results.json. This is
a Node.js numerical replay of browser code, not a new browser training run.
The independent Python tensor script was also rerun without PyTorch; it matches
the retained digit-1 tensors. NumPy/runtime versions are recorded in S1 follow-up.

The retained 10000 MNIST predictions were audited by recomputing correct count,
confusion matrix and macro-F1. New MNIST inference was not run. Original metrics
remain 197/200 synthetic glyphs and 5773/10000 authentic MNIST images.

No PyTorch checkpoint, export_viz.py or PyTorch-to-browser export comparison is
claimed or fabricated. A separate PyTorch experiment would require its own
versioned source, weights, export script and numerical comparison evidence.
''',encoding='utf-8')
    summary=S4/'software validation summary.json';old=read_json(summary)
    old['historical_scope']='Original results dated 17 October 2026; follow-up evidence is separately dated and versioned.'
    old['follow_up_2026_10_18']={'portal_real_http_checks':17,'retention_checks':9,'backend_tests':69,'frontend_tests':142,
        'failed':0,'provider_mode':'simulated external provider; real portal/gateway transport',
        'source_directory':'../S1 configuration and source/Portal integration revision 2026-10-18',
        'portal_results':'Portal integration validation 2026-10-18/results.json','retention_results':'Retention validation 2026-10-18/results.json',
        'production_provider_retention':'NOT_VERIFIED','production_deletion_and_redis_ttl':'NOT_VERIFIED'}
    write_json(summary,old)
    readme=MATERIALS/'README.txt';text=readme.read_text(encoding='utf-8')
    text=text.replace('Technical validation performed on 17 October 2026','Technical validation performed on 17 October 2026; integration/retention follow-up on 18 October 2026')
    text=text.replace('MAIN RESULTS\n','MAIN RESULTS (ORIGINAL 17 OCTOBER RUN)\n')
    text=text.replace('The new public portal has a tutor extension point; its complete authenticated bridge to the AI gateway remains an integration task.',
        'The frozen 17 October public portal has a tutor extension point. The current authenticated bridge is covered separately by the 18 October integration follow-up with a simulated external provider.')
    text+='''\n18 OCTOBER FOLLOW-UP\nThe original source archives and experiment results are preserved. New source\narchives and per-file hashes are in S1/Portal integration revision 2026-10-18.\nS4/Portal integration validation 2026-10-18 contains 17 passing browser/HTTP\nchecks, raw synthetic requests/records, screenshots, 69 backend and 142 frontend\nregression tests, successful builds/typecheck and the current validated domain schema.\nOnly the external provider is simulated. No new live-provider RAG scores are claimed.\nS4/Retention validation 2026-10-18 contains nine passing SQLite/PostgreSQL and\nRedis checks and explicit NOT_VERIFIED production/provider-account items.\nS3/Model verification 2026-10-18 documents the browser ten-class model selection,\n200-image numerical replay, independent Python replay and retained MNIST metric audit.\n\nReproduction scripts are in validation/ in the new backend source archive.\nFor the model replay: node validation/cnn_replay.cjs <MATERIALS_ROOT> <NEW_OUTPUT>.\nFor the independent tensor replay: python "check CNN tensors.py" --output <NEW_JSON>.\nNumPy is required. The original tensor report is preserved when --output is supplied.\nThe article DOCX was not edited. See REVIEW RESPONSES 2026-10-18.txt for the four\nreviewer items and remaining operational confirmations.\n'''
    readme.write_text(text,encoding='utf-8')
    from docx import Document
    document=Document(MATERIALS/'AIEdu Supplementary Materials.docx')
    for p in document.paragraphs:
        if 'The current public portal provides a tutor extension point;' in p.text:
            p.text=p.text.replace('The current public portal provides a tutor extension point; a production identity and course-authorization bridge remains an integration task.',
                'The frozen 17 October portal contains a tutor extension point. The 18 October addendum separately evaluates the current authenticated bridge using a simulated external provider.')
    document.add_heading('Addendum: integration and retention validation, 18 October 2026',level=1)
    document.add_paragraph('The original 17 October source snapshots and experiment results above are preserved. The follow-up uses separately identified current source archives in S1/Portal integration revision 2026-10-18. It does not replace or rerun the original live-provider RAG experiment.')
    table=document.add_table(rows=1,cols=2);table.rows[0].cells[0].text='Follow-up check';table.rows[0].cells[1].text='Observed result'
    for label,value in [('Real portal/gateway browser and HTTP checks','17 passed; external provider simulated'),('Current backend regression tests','69 passed'),('Current frontend regression tests','142 passed'),('Frontend and AI builds; AI typecheck; domain schema validation','Passed'),('SQLite/PostgreSQL cleanup and Redis TTL checks','9 passed, isolated local instances'),('Frozen browser classifier replay','200/200 predictions match retained predictions; 197 correct'),('Independent Python tensor replay','All retained example tensors match; PyTorch not used'),('Provider account and production retention/Redis configuration','NOT_VERIFIED')]:
        cells=table.add_row().cells;cells[0].text=label;cells[1].text=value
    for paragraph in [
        'Two authenticated browser contexts exchanged three turns concurrently. Recorded provider timestamps establish actual overlap. Gateway metadata and prompts retain the correct user, course and lesson; unauthorized lecture requests are denied before the gateway. Six controlled failure cases verify error presentation and prevent incomplete assistant answers from being stored as completed.',
        'The gateway now sends store=false explicitly, uses configurable provider timeouts/retries, and rejects a provider stream without response.completed. Requests use a local deterministic HTTP fixture, so no live-provider connectivity or model-quality conclusion follows from this integration run.',
        'Local cleanup tests use a 30-day retention setting and a 25-request maintenance interval. Expired records remain until maintenance is invoked. Redis default TTLs are observed on actual keys; a separate one-second override checks physical expiry. The isolated Redis test server is version 5.0.14.1, not the production Compose Redis 7 service. PostgreSQL retention tables are unchanged; the optional vector column/index is omitted because this host lacks pgvector.',
        'Account-level provider retention, deployed deletion operation, backups and production Redis TTLs remain unverified. The store flag and local cleanup test do not establish these operational properties.',
        'The selected evaluated laboratory is the JavaScript browser ten-class model. The 18 October replay uses the original source modules, checkpoint and 200 pixel arrays without retraining. MNIST aggregate metrics are recalculated from retained predictions; new MNIST inference is not claimed. Independent Python arithmetic is not a PyTorch export verification.']:
        document.add_paragraph(paragraph)
    document.save(MATERIALS/'AIEdu Supplementary Materials.docx')
    (MATERIALS/'REVIEW RESPONSES 2026-10-18.txt').write_text('''Ответы на четыре красных замечания в статье 17_10

1. Передача пользователя и разрешённой лекции реализована и проверена локально
через настоящий браузер, Django и Next.js. Только внешний провайдер имитируется.
Проверены серверные user_id/course_id/lesson_id, текст лекции, отказ чужой лекции
и игнорирование подставленных клиентом полей. S4: Portal integration validation
2026-10-18/results.json, gateway records.json, provider requests.json, screenshots.
Новая версия исходников включена отдельно в S1; старые архивы сохранены.

2. Два пользователя обменялись тремя сообщениями каждый одновременно. Проверены
отдельные браузерные истории, запросы к провайдеру и записи шлюза; временные
метки подтверждают перекрытие запросов. Проверены шесть сценариев сбоев:
retrieval, HTTP-отказ провайдера, timeout, EOF без завершения, зависание потока,
повреждённый SSE. Все 17 проверок прошли. Это контролируемые интеграционные
испытания с имитатором провайдера, а не испытания реального OpenAI.

3. Шлюз явно передаёт store=false. На временных SQLite/PostgreSQL проверено
удаление просроченных записей после 25-го вызова обслуживания и сохранение
свежих записей. В отдельном Redis проверены TTL пяти видов ключей и реальное
истечение при ускоренном TTL. Все девять проверок прошли. Фактические настройки
аккаунта провайдера, удаления на рабочем сервере и рабочего Redis НЕ ПОДТВЕРЖДЕНЫ.
Список необходимых подтверждений: S4/Retention validation 2026-10-18/deployment
confirmation status.json. Очистка вызывается активностью, а не по расписанию.

4. Материалы описывают браузерную модель на 10 классов без утверждения о PyTorch.
Повторный расчёт всех 200 сохранённых изображений воспроизвёл классы; независимый
Python-расчёт воспроизвёл тензоры. Статистика 10000 MNIST-предсказаний пересчитана
из сохранённых результатов; новый inference MNIST не запускался. Source/checkpoint
PyTorch и export_viz.py не добавлялись: соответствующего эксперимента в пакете нет.
См. S3/Model verification 2026-10-18/Model scope.txt и replay results.json.

Дополнительно: 69 backend-тестов, 142 frontend-теста, обе production-сборки,
typecheck AI и генерация/валидация актуальной domain OpenAPI прошли.
Статья 17_10 не редактировалась. Утверждать проверенные рабочие настройки
хранения или реальный провайдер по этим локальным испытаниям нельзя.
''',encoding='utf-8')
    changed_code=['AI-edu-backend/TUTOR_INTEGRATION.md','AI-edu-backend/services/ai-tutor/.env.example',
        'AI-edu-backend/services/ai-tutor/app/api/chat/route.ts']
    changed_code.extend(p.relative_to(ROOT).as_posix() for p in sorted((ROOT/'AI-edu-backend/validation').glob('*')) if p.is_file())
    changed_materials=['AIEdu Supplementary Materials.docx','README.txt','check CNN tensors.py','REVIEW RESPONSES 2026-10-18.txt',
        'supplementary/S1 configuration and source/repository provenance.json','supplementary/S4 API and validation/software validation summary.json',
        'SHA256SUMS.txt','SHA256SUMS Publication.txt']
    changed_materials.extend(p.relative_to(MATERIALS).as_posix() for folder in (S1,S3,PORTAL,RETENTION) for p in sorted(folder.rglob('*')) if p.is_file())
    write_json(MATERIALS/'CHANGED FILES 2026-10-18.json',{'code_files':changed_code,'material_files':sorted(changed_materials+['CHANGED FILES 2026-10-18.json']),
        'article_modified':False,'frontend_source_modified':False,'final_archive':'../AIEdu Supplementary Materials.zip'})
    # All edited/new material files are covered. Exclude checksum manifests themselves.
    files=sorted(p for p in MATERIALS.rglob('*') if p.is_file() and '__pycache__' not in p.parts and p.suffix!='.pyc' and p.name not in ['SHA256SUMS.txt','SHA256SUMS Publication.txt'])
    for filename,publication in [('SHA256SUMS.txt',False),('SHA256SUMS Publication.txt',True)]:
        selected=[p for p in files if publication or 'S5 Publication Figures' not in p.parts]
        (MATERIALS/filename).write_text(''.join(sha(p)+'  '+p.relative_to(MATERIALS).as_posix()+'\n' for p in selected),encoding='utf-8')
    archive=ROOT/'AIEdu Supplementary Materials.zip'
    if archive.exists():raise FileExistsError('Refusing to replace an existing final archive')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(MATERIALS.rglob('*')):
            if p.is_file() and '__pycache__' not in p.parts and p.suffix!='.pyc':z.write(p,p.relative_to(MATERIALS).as_posix())
    print(json.dumps({'code_files':changed_code,'new_material_files':len(changed_materials),'archive':str(archive),'archive_sha256':sha(archive)},ensure_ascii=False,indent=2))


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
